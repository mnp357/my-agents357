from liquidation_report import analyse,number,export
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook

def row(symbol='XRP',side='Long',liq='93'):
    return dict(symbol=symbol,side=side,price='$100',liquidation=liq,volume='$1.20M')
data={'eligible_symbols':['XRP'],'coins':[{'symbol':'XRP','error':None,'pages':[{'captured_at':'2026-09-25T00:00:00Z','rows':[
    row(),row(liq='92.999'),row(side='Short',liq='107'),row(side='Short',liq='107.001'),row(symbol='xyz:GOLD'),row(liq='—'),row(liq='105')]}]}]}
selected,coverage=analyse(data)
assert len(selected)==2,selected
assert [r[1] for r in selected]==['Лонг','Шорт']
assert all(r[5]==1200000 for r in selected)
assert coverage[0][3]==1
assert number('NaN') is None
print('7% boundaries, direction, XYZ substring, missing values and amount parsing: passed')

with TemporaryDirectory() as directory:
    sample={'source':'CoinGlass','started_at':'2026-10-08T00:00:00Z',
            'finished_at':'2026-10-08T00:01:00Z','eligible_symbols':['BTC'],
            'coins':[{'symbol':'BTC','error':None,'pages':[{'captured_at':'2026-10-08T00:00:30Z','rows':[
                row('BTC','Long','94') | {'volume':'$10M'},
                row('BTC','Long','99') | {'volume':'$2M'},
                row('BTC','Short','106') | {'volume':'$8M'},
                row('BTC','Short','101') | {'volume':'$3M'},
            ]}]}]}
    file=Path(directory)/'report.xlsx'
    export(sample,file)
    sheet=load_workbook(file,read_only=True,data_only=True)['Сводка']
    assert [sheet.cell(10,col).value for col in (4,5,6)] == [
        'Ближайший лонг','Ближайший шорт','Позиций в пределах 7%']
    assert [sheet.cell(11,col).value for col in (4,5,6)] == [99,101,4]
    assert all(sheet.cell(row,col).alignment.horizontal == 'center'
               for row in (10,11) for col in range(1,7))
print('Summary nearest levels and centered alignment: passed')
