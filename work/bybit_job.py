"""Collect the available Bybit account history and send an auditable workbook."""
import fcntl
import json
import os
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from bybit_client import BybitError, Client, pages
from telegram_bridge import send, state

ROOT = Path('/var/lib/my-agents')
STATUS = ROOT / 'bybit-status.json'
CANCEL = ROOT / 'bybit-cancel'
KEYS = ROOT / 'bybit-api.json'
DAY = 86400000


def update(**fields):
    fields['updated_at'] = datetime.now(timezone.utc).isoformat()
    temp = STATUS.with_suffix('.tmp')
    temp.write_text(json.dumps(fields, ensure_ascii=False))
    os.replace(temp, STATUS)


def decimal(value):
    try:
        return Decimal(str(value or 0))
    except InvalidOperation:
        return Decimal(0)


def utc(ms):
    try:
        return datetime.fromtimestamp(int(ms) / 1000, timezone.utc).strftime('%Y-%m-%d %H:%M:%S')
    except (ValueError, TypeError, OverflowError):
        return ''


def windows(start, end, days):
    width = days * DAY - 1000
    cursor = start
    while cursor <= end:
        stop = min(cursor + width, end)
        yield cursor, stop
        cursor = stop + 1000


def collect_series(client, name, path, list_key, start, end, days, extra, coverage):
    rows = []
    intervals = list(windows(start, end, days))
    for index, (lo, hi) in enumerate(intervals, 1):
        if CANCEL.exists():
            coverage.append((name, utc(lo), utc(hi), 'Остановлено', 0, ''))
            break
        update(state='running', message=f'Bybit: {name}, период {index}/{len(intervals)}', source=name)
        params = dict(extra, startTime=lo, endTime=hi)
        try:
            batch = list(pages(client, path, params, list_key, CANCEL.exists))
            rows.extend(batch)
            coverage.append((name, utc(lo), utc(hi), 'Получено', len(batch), ''))
        except BybitError as exc:
            coverage.append((name, utc(lo), utc(hi), 'Ошибка', 0, str(exc)))
    return rows


def collect_funding(client, start, end, coverage):
    rows = []
    intervals = list(windows(start, end, 7))
    for index, (lo, hi) in enumerate(intervals, 1):
        if CANCEL.exists():
            coverage.append(('Финансовый счёт', utc(lo), utc(hi), 'Остановлено', 0, ''))
            break
        update(state='running', message=f'Bybit: финансовый счёт, период {index}/{len(intervals)}')
        try:
            batch = list(pages(client, '/v5/asset/fundinghistory',
                               {'createTimeFrom': lo // 1000, 'createTimeTo': hi // 1000,
                                'limit': 100}, 'list', CANCEL.exists))
            rows.extend(batch)
            coverage.append(('Финансовый счёт', utc(lo), utc(hi), 'Получено', len(batch), ''))
        except BybitError as exc:
            coverage.append(('Финансовый счёт', utc(lo), utc(hi), 'Ошибка', 0, str(exc)))
    return rows


def sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        # Prevent spreadsheet formula injection from exchange-provided text.
        ws.append([("'" + x if isinstance(x, str) and x.startswith(('=', '+', '-', '@')) else x)
                   for x in row])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='173F3D')
        cell.font = Font(color='FFFFFF', bold=True)
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = min(45, max(12, max(len(str(c.value or '')) for c in list(column)[:300]) + 2))
        ws.column_dimensions[letter].width = width
    return ws


def export(path, data, period_start, period_end, uid):
    wb = Workbook()
    wb.remove(wb.active)
    closed = data['closed']
    executions = data['executions']
    ledger = data['ledger']
    deposits = data['deposits']
    internal_deposits = data['internal_deposits']
    withdrawals = data['withdrawals']
    funding = data['funding']
    transfers = data['transfers']
    coverage = data['coverage']
    pnl = sum((decimal(x.get('closedPnl')) for x in closed), Decimal(0))
    grouped = defaultdict(lambda: [0, Decimal(0)])
    monthly = defaultdict(lambda: [0, Decimal(0), 0])
    for row in closed:
        group = (row.get('_category', ''), row.get('symbol', ''))
        grouped[group][0] += 1
        grouped[group][1] += decimal(row.get('closedPnl'))
        month = utc(row.get('updatedTime'))[:7]
        monthly[month][0] += 1
        monthly[month][1] += decimal(row.get('closedPnl'))
        monthly[month][2] += decimal(row.get('closedPnl')) > 0
    spot_grouped = defaultdict(lambda: [0, Decimal(0), Decimal(0)])
    for row in executions:
        if row.get('_category') != 'spot':
            continue
        stat = spot_grouped[row.get('symbol', '')]
        stat[0] += 1
        if row.get('side') == 'Buy':
            stat[1] += decimal(row.get('execValue'))
        elif row.get('side') == 'Sell':
            stat[2] += decimal(row.get('execValue'))
    movements = defaultdict(lambda: [Decimal(0), Decimal(0), 0, 0])
    for row in deposits:
        if str(row.get('status', '')).lower() in ('3', 'success'):
            stat = movements[row.get('coin', '')]
            stat[0] += decimal(row.get('amount'))
            stat[2] += 1
    for row in internal_deposits:
        if str(row.get('status', '')).lower() in ('2', 'success'):
            stat = movements[row.get('coin', '')]
            stat[0] += decimal(row.get('amount'))
            stat[2] += 1
    for row in withdrawals:
        if str(row.get('status', '')).lower() in ('success', 'completed'):
            stat = movements[row.get('coin', '')]
            stat[1] += decimal(row.get('amount'))
            stat[3] += 1
    failed = sum(1 for x in coverage if x[3] != 'Получено')
    overview = [
        ('Период, UTC', f'{utc(period_start)} — {utc(period_end)}'),
        ('UID', uid),
        ('Закрытые фьючерсные записи', len(closed)),
        ('PnL закрытых фьючерсных записей', float(pnl)),
        ('Спотовые исполнения', sum(x.get('_category') == 'spot' for x in executions)),
        ('Прочие исполнения', sum(x.get('_category') != 'spot' for x in executions)),
        ('Записи журнала Unified', len(ledger)),
        ('Входящие депозиты', len(deposits)),
        ('Внутренние поступления', len(internal_deposits)),
        ('Исходящие выводы', len(withdrawals)),
        ('Внутренние переводы', len(transfers)),
        ('Записи финансового счёта', len(funding)),
        ('Периоды с ошибкой/остановкой', failed),
        ('Статус полноты', 'НЕПОЛНЫЙ' if failed else 'Доступные API-периоды собраны; данные старше лимитов API не включены'),
        ('Методика', 'Пополнения, выводы и внутренние переводы не включены в торговый PnL.'),
        ('Спотовый PnL', 'Не рассчитывается без подтверждённой начальной себестоимости и полной истории.'),
        ('Ограничение', 'История Unified API: до 2 лет. Старые и до-UTA данные могут требовать отдельной выгрузки Bybit.'),
    ]
    sheet(wb, 'Сводка', ['Показатель', 'Значение'], overview)
    sheet(wb, 'Фьючерсы по монетам', ['Категория', 'Монета', 'Закрытий', 'PnL'],
          [(a, b, n, float(v)) for (a, b), (n, v) in sorted(grouped.items())])
    sheet(wb, 'Фьючерсы по месяцам', ['Месяц UTC', 'Закрытий', 'PnL', 'Прибыльных закрытий'],
          [(month, n, float(value), wins) for month, (n, value, wins) in sorted(monthly.items())])
    sheet(wb, 'Спот по монетам', ['Монета', 'Исполнений', 'Покупки в валюте котировки', 'Продажи в валюте котировки', 'Примечание'],
          [(symbol, n, float(buys), float(sells), 'Не является прибылью')
           for symbol, (n, buys, sells) in sorted(spot_grouped.items())])
    sheet(wb, 'Движения по валютам', ['Валюта', 'Поступило', 'Выведено', 'Поступлений', 'Выводов'],
          [(coin, float(inp), float(out), ni, no) for coin, (inp, out, ni, no) in sorted(movements.items())])
    sheet(wb, 'Закрытые фьючерсы', ['Категория', 'Монета', 'Сторона', 'Время закрытия UTC', 'PnL', 'Объём', 'Средняя цена входа', 'Средняя цена выхода', 'Комиссия открытия', 'Комиссия закрытия', 'ID ордера'],
          [(x.get('_category'), x.get('symbol'), x.get('side'), utc(x.get('updatedTime')),
            float(decimal(x.get('closedPnl'))), float(decimal(x.get('qty'))),
            x.get('avgEntryPrice'), x.get('avgExitPrice'), float(decimal(x.get('openFee'))),
            float(decimal(x.get('closeFee'))), x.get('orderId')) for x in closed])
    sheet(wb, 'Исполнения сделок', ['Рынок', 'Время UTC', 'Монета', 'Сторона', 'Количество', 'Цена', 'Сумма', 'Комиссия', 'Валюта комиссии', 'ID исполнения', 'ID ордера'],
          [(x.get('_category'), utc(x.get('execTime')), x.get('symbol'), x.get('side'),
            float(decimal(x.get('execQty'))), x.get('execPrice'), x.get('execValue'),
            x.get('execFee'), x.get('feeCurrency'), x.get('execId'), x.get('orderId')) for x in executions])
    sheet(wb, 'Журнал Unified', ['Время UTC', 'Рынок', 'Тип', 'Монета', 'Валюта', 'Денежный поток', 'Funding', 'Комиссия', 'Изменение баланса', 'Баланс', 'ID'],
          [(utc(x.get('transactionTime')), x.get('category'), x.get('type'), x.get('symbol'), x.get('currency'),
            x.get('cashFlow'), x.get('funding'), x.get('fee'), x.get('change'), x.get('cashBalance'), x.get('id')) for x in ledger])
    sheet(wb, 'Поступления', ['Время UTC', 'Монета', 'Сумма', 'Сеть', 'Статус', 'Тип', 'ID', 'TxID'],
          [(utc(x.get('successAt')), x.get('coin'), x.get('amount'), x.get('chain'), x.get('status'),
            'Блокчейн', x.get('id'), x.get('txID')) for x in deposits] +
          [(utc(int(x.get('createdTime', 0)) * 1000), x.get('coin'), x.get('amount'), '', x.get('status'),
            'Внутри Bybit', x.get('id'), x.get('txID')) for x in internal_deposits])
    sheet(wb, 'Выводы', ['Время UTC', 'Монета', 'Сумма', 'Комиссия', 'Сеть', 'Статус', 'Тип', 'ID', 'TxID'],
          [(utc(x.get('createTime')), x.get('coin'), x.get('amount'), x.get('withdrawFee'), x.get('chain'),
            x.get('status'), x.get('withdrawType'), x.get('withdrawId'), x.get('txID')) for x in withdrawals])
    sheet(wb, 'Внутренние переводы', ['Время UTC', 'Монета', 'Сумма', 'Откуда', 'Куда', 'Статус', 'ID'],
          [(utc(x.get('timestamp')), x.get('coin'), x.get('amount'), x.get('fromAccountType'),
            x.get('toAccountType'), x.get('status'), x.get('transferId')) for x in transfers])
    sheet(wb, 'Финансовый счёт', ['Время UTC', 'Монета', 'Направление', 'Сумма', 'Тип', 'Описание', 'ID'],
          [(utc(int(x.get('createTime', 0)) * 1000), x.get('currency'), x.get('ioDirection'), x.get('txnAmt'),
            x.get('showBusiTypeEn'), x.get('descriptionEn'), x.get('currcCursor')) for x in funding])
    sheet(wb, 'Охват и ошибки', ['Источник', 'Начало UTC', 'Конец UTC', 'Статус', 'Записей', 'Причина'], coverage)
    wb.save(path)
    return {'closed': len(closed), 'spot': sum(x.get('_category') == 'spot' for x in executions),
            'deposits': len(deposits) + len(internal_deposits), 'withdrawals': len(withdrawals), 'failed_intervals': failed}


def run():
    if not KEYS.exists():
        update(state='failed', message='Сначала подключите ключ Bybit на сайте.')
        return
    credentials = json.loads(KEYS.read_text())
    client = Client(credentials['key'], credentials['secret'], credentials['region'])
    info = client.verify_read_only()
    end = int(datetime.now(timezone.utc).timestamp() * 1000)
    start = int((datetime.now(timezone.utc) - timedelta(days=730)).timestamp() * 1000)
    data = {k: [] for k in ('closed', 'executions', 'ledger', 'deposits', 'internal_deposits', 'withdrawals', 'funding', 'transfers', 'coverage')}
    cov = data['coverage']
    specs = [
        ('Журнал Unified', 'ledger', '/v5/account/transaction-log', 'list', 7, {'limit': 50}),
        ('Спот', 'executions', '/v5/execution/list', 'list', 7, {'category': 'spot', 'limit': 100}),
        ('Линейные исполнения', 'executions', '/v5/execution/list', 'list', 7, {'category': 'linear', 'limit': 100}),
        ('Инверсные исполнения', 'executions', '/v5/execution/list', 'list', 7, {'category': 'inverse', 'limit': 100}),
        ('Опционные исполнения', 'executions', '/v5/execution/list', 'list', 7, {'category': 'option', 'limit': 100}),
        ('Линейный PnL', 'closed', '/v5/position/closed-pnl', 'list', 7, {'category': 'linear', 'limit': 100}),
        ('Инверсный PnL', 'closed', '/v5/position/closed-pnl', 'list', 7, {'category': 'inverse', 'limit': 100}),
        ('Депозиты', 'deposits', '/v5/asset/deposit/query-record', 'rows', 29, {'limit': 50}),
        ('Внутренние поступления', 'internal_deposits', '/v5/asset/deposit/query-internal-record', 'rows', 29, {'limit': 50}),
        ('Выводы', 'withdrawals', '/v5/asset/withdraw/query-record', 'rows', 29, {'limit': 50, 'withdrawType': 2}),
        ('Внутренние переводы', 'transfers', '/v5/asset/transfer/query-inter-transfer-list', 'list', 7, {'limit': 50}),
    ]
    for name, dest, path, key, days, extra in specs:
        if CANCEL.exists():
            break
        result = collect_series(client, name, path, key, start, end, days, extra, cov)
        if dest in ('closed', 'executions'):
            for row in result:
                row['_category'] = extra['category']
        data[dest].extend(result)
    if not CANCEL.exists():
        data['funding'] = collect_funding(client, start, end, cov)
    if not any(data[k] for k in data if k != 'coverage') and any(x[3] == 'Ошибка' for x in cov):
        raise BybitError('Bybit не вернул данные. Проверьте права ключа и лист ошибок после исправления.')
    reports = ROOT / 'reports'
    reports.mkdir(exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')
    file = reports / ('Bybit-' + stamp + '.xlsx')
    summary = export(file, data, start, end, info['uid'])
    current = state('bybit')
    caption = (f"Анализ Bybit готов. Спот: {summary['spot']} исполнений; закрытые фьючерсы: "
               f"{summary['closed']}; поступления: {summary['deposits']}; выводы: {summary['withdrawals']}.")
    if summary['failed_intervals']:
        caption += ' Есть пропуски: лист «Охват и ошибки».'
    if current and current['owner_chat_id'] is not None:
        from agent_job import send_document
        send_document(current['token'], current['owner_chat_id'], file, caption)
    update(state='cancelled' if CANCEL.exists() else 'completed', message=caption,
           file=file.name, **summary)


if __name__ == '__main__':
    with open(ROOT / 'bybit.lock', 'a') as lock:
        try:
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            raise SystemExit(0)
        try:
            run()
        except Exception:
            update(state='failed', message='Анализ Bybit не завершён. Проверьте ключ и повторите запуск.')
            current = state('bybit')
            if current and current['owner_chat_id'] is not None:
                try:
                    send(current['token'], current['owner_chat_id'], 'Анализ Bybit не завершён. Проверьте подключение на сайте.')
                except Exception:
                    pass
