import json
import os
from pathlib import Path
from openpyxl import load_workbook
from liquidation_report import export

root=Path('/var/lib/my-agents')
status_path=root/'job-status.json'
status=json.loads(status_path.read_text())
assert status['state']=='completed', 'Wait for active report to finish'
old=root/'reports'/status['file']
stamp=old.stem.removeprefix('Hyperliquid-')
raw=root/'reports'/(stamp+'.json')
data=json.loads(raw.read_text())
new=old.with_name(old.stem+'-spaced.xlsx')
export(data,new)
before=load_workbook(old,data_only=True)
after=load_workbook(new,data_only=True)
old_rows=list(before['Позиции'].values)
new_rows=list(after['Позиции'].values)
assert [r for r in new_rows if any(v is not None for v in r)]==old_rows
blank=[]
for i,row in enumerate(new_rows):
    if all(v is None for v in row):
        assert 1<i<len(new_rows)-1
        assert new_rows[i-1][0] and new_rows[i+1][0]
        assert new_rows[i-1][0]!=new_rows[i+1][0]
        blank.append(i+1)
assert len(blank)==len({r[0] for r in old_rows[1:]})-1
for name in ('Сводка','Охват'):
    assert list(before[name].values)==list(after[name].values)
status['file']=new.name
temporary=status_path.with_suffix('.tmp')
temporary.write_text(json.dumps(status,ensure_ascii=False))
os.replace(temporary,status_path)
print(json.dumps({'file':new.name,'positions':len(old_rows)-1,'blank_rows':blank},ensure_ascii=False))
