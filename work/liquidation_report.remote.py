"""Excel export for the deployed Linux worker (standalone runtime)."""
from decimal import Decimal, InvalidOperation
from collections import defaultdict

def number(text):
    text=str(text).replace('$','').replace(',','').replace('\u00a0','').replace(' ','').strip()
    multiplier=Decimal(1)
    if text[-1:] in ('K','M','B'):
        multiplier={'K':Decimal(1000),'M':Decimal(1000000),'B':Decimal(1000000000)}[text[-1]]
        text=text[:-1]
    try:
        value=Decimal(text)*multiplier
        return value if value.is_finite() else None
    except InvalidOperation:
        return None

def analyse(data):
    selected=[]
    coverage=[]
    for coin in data['coins']:
        missing=0
        for page in coin['pages']:
            for row in page['rows']:
                if 'XYZ' in row['symbol'].upper(): continue
                price,liq,volume=(number(row.get(k,'')) for k in ('price','liquidation','volume'))
                if price is None or price<=0 or liq is None or liq<=0 or volume is None or volume<0:
                    missing+=1
                    continue
                distance=(liq-price)/price
                eligible=(row['side']=='Long' and -Decimal('.07')<=distance<=0) or (row['side']=='Short' and 0<=distance<=Decimal('.07'))
                if eligible:
                    selected.append([row['symbol'],'Лонг' if row['side']=='Long' else 'Шорт',float(price),float(liq),float(distance),float(volume),page['captured_at']])
        coverage.append([coin['symbol'],len(coin['pages']),sum(len(p['rows']) for p in coin['pages']),missing,'Ошибка' if coin['error'] else 'Проверено',coin['error'] or ''])
    done={c['symbol'] for c in data['coins']}
    coverage.extend([[s,0,0,0,'Не проверено',''] for s in data['eligible_symbols'] if s not in done])
    selected.sort(key=lambda r:(r[0],r[1],abs(r[4])))
    return selected,coverage

def export(data,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    rows,coverage=analyse(data)
    wb=Workbook()
    ws=wb.active
    ws.title='Сводка'
    ws.append(['Hyperliquid — ликвидации в пределах 7%'])
    ws.append(['Источник',data['source']])
    ws.append(['Начало (UTC)',data['started_at']])
    ws.append(['Окончание (UTC)',data.get('finished_at','')])
    ws.append(['Охват','До 3 страниц каждой монеты, позиции по убыванию размера; XYZ исключён'])
    ws.append(['Ограничение','Последовательная выборка публичной таблицы; суммы и цены округлены CoinGlass'])
    ws.append(['Примечание','Объём позиции не равен гарантированному объёму ликвидации. Пустые уровни пропущены.'])
    ws.append(['Результат','Неполный сбор' if any(c[4]!='Проверено' for c in coverage) or data.get('cancelled') else 'Все выбранные монеты проверены в пределах 3 страниц'])
    ws.append([])
    ws.append(['Монета','Лонги, USD','Шорты, USD','Ближайший лонг','Ближайший шорт','Позиций в пределах 7%'])
    groups=defaultdict(list)
    for row in rows: groups[row[0]].append(row)
    for symbol in sorted(groups):
        long=[r for r in groups[symbol] if r[1]=='Лонг']
        short=[r for r in groups[symbol] if r[1]=='Шорт']
        ws.append([symbol,sum(r[5] for r in long),sum(r[5] for r in short),min(long,key=lambda r:abs(r[4]))[3] if long else None,min(short,key=lambda r:abs(r[4]))[3] if short else None,len(groups[symbol])])
    detail=wb.create_sheet('Позиции')
    detail.append(['Монета','Сторона','Текущая цена, USD','Ликвидация, USD','Расстояние, %','Размер позиции, USD','Сбор (UTC)'])
    for row in rows: detail.append(row)
    cov=wb.create_sheet('Охват')
    cov.append(['Монета','Страниц','Строк','Нет корректного уровня/числа','Статус','Ошибка'])
    for row in coverage: cov.append(row)
    for sheet,header in ((ws,10),(detail,1),(cov,1)):
        sheet.freeze_panes=f'A{header+1}'
        sheet.auto_filter.ref=f'A{header}:{get_column_letter(sheet.max_column)}{max(header,sheet.max_row)}'
        sheet.sheet_view.showGridLines=False
        for cell in sheet[header]:
            cell.fill=PatternFill('solid',fgColor='18384A')
            cell.font=Font(color='FFFFFF',bold=True)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[header].height=32
        for col in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(col)].width=23
    ws.column_dimensions['B'].width=34
    ws.row_dimensions[1].height=30
    ws['A1'].font=Font(size=18,bold=True,color='18384A')
    for n in range(2,9): ws.merge_cells(start_row=n,start_column=2,end_row=n,end_column=6)
    for row in detail.iter_rows(min_row=2):
        for cell in row[2:4]: cell.number_format='#,##0.########'
        row[4].number_format='0.00%;[Red]-0.00%'
        row[5].number_format='#,##0.00'
    detail.column_dimensions['G'].width=35
    cov.column_dimensions['D'].width=30
    cov.column_dimensions['F'].width=65
    for row in ws.iter_rows(min_row=11):
        for cell in row[1:5]: cell.number_format='#,##0.########'
    wb.save(path)
    return {'positions':len(rows),'coins':len(groups),'checked':sum(c[4]=='Проверено' for c in coverage),'total':len(coverage),'errors':sum(c[4]!='Проверено' for c in coverage)}
